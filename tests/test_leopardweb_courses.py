"""Tests for the calendar feed reader.

The feed itself decides the row grain, so these tests pin the two things this
script is responsible for: reading the right report for the grain the user
asked for, and mapping feed columns onto spreadsheet columns without losing
numbers or inventing values.
"""

import csv
import json

import pytest

import leopardweb_courses as lw


SECTION_CSV = (
    "term_uid,term,crn,subject,course_number,section_number,title,"
    "schedule_type,status,credit_hours,faculty,seats_capacity,seats_available,"
    "enrollment_current,meeting_days,meeting_times,meeting_type,room,"
    "room_capacity,meeting_count\r\n"
    "202710,Fall 2026,16036,Mathematics (MATH),2300,3,Discrete Structures,"
    "lecture,active,4,Mark Mixer,30,5,25,TR,10:00-11:45,lecture,WENTW 214,"
    "40,2\r\n"
    "202710,Fall 2026,17309,Mathematics (MATH),1525,1A,Calculus,"
    "lecture,active,4,Mark Mixer,24,0,24,MW,08:00-09:15,lecture,WENTW 206,"
    ",2\r\n"
)

MEETING_CSV = (
    "term_uid,term,crn,subject,course_number,section_number,title,"
    "schedule_type,status,credit_hours,faculty,seats_capacity,seats_available,"
    "enrollment_current,day,day_of_week,begin_time,end_time,meeting_type,"
    "building,building_name,room_number,room,room_capacity\r\n"
    "202710,Fall 2026,16036,Mathematics (MATH),2300,3,Discrete Structures,"
    "lecture,active,4,Mark Mixer,30,5,25,tuesday,2,10:00,11:45,lecture,"
    "WENTW,Wentworth Hall,214,WENTW 214,40\r\n"
    "202710,Fall 2026,16036,Mathematics (MATH),2300,3,Discrete Structures,"
    "lecture,active,4,Mark Mixer,30,5,25,thursday,4,10:00,11:45,lecture,"
    "WENTW,Wentworth Hall,214,WENTW 214,40\r\n"
)


class FakeResponse:
    def __init__(self, text, content_type="text/csv; charset=utf-8"):
        self.text = text
        self.headers = {"Content-Type": content_type}

    def raise_for_status(self):
        return None


class FakeSession:
    """Records the paths asked for and answers with canned CSV bodies."""

    def __init__(self, bodies):
        self.bodies = bodies
        self.calls = []
        self.headers = {}

    def get(self, url, params=None, timeout=None):
        self.calls.append((url, params))
        for path, body in self.bodies.items():
            if url.endswith(path):
                return FakeResponse(body)
        raise AssertionError(f"unexpected URL {url}")


@pytest.fixture
def client():
    feed = lw.CalendarFeedClient(base_url="https://calendar.example")
    feed.session = FakeSession({
        "/reports/sections": SECTION_CSV,
        "/reports/meeting_times": MEETING_CSV,
    })
    return feed


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


class TestReportSelection:
    def test_sections_reads_the_sections_report(self, client):
        rows = client.get_sections("202710")

        assert client.session.calls[0][0].endswith("/reports/sections")
        assert client.session.calls[0][1] == {"term_uid": "202710"}
        assert len(rows) == 2

    def test_meeting_times_reads_the_meeting_times_report(self, client):
        rows = client.get_meeting_times("202710")

        assert client.session.calls[0][0].endswith("/reports/meeting_times")
        assert len(rows) == 2

    def test_a_term_with_no_data_names_list_terms(self, client):
        client.session.bodies["/reports/sections"] = "term_uid,term,crn\r\n"

        with pytest.raises(lw.CalendarFeedError, match="--list-terms"):
            client.get_sections("999999")

    def test_a_missing_report_says_the_server_is_old(self, client):
        class Missing(FakeResponse):
            def __init__(self):
                super().__init__("")
                self.status_code = 404

            def raise_for_status(self):
                raise lw.requests.HTTPError("404 Not Found", response=self)

        client.session.get = lambda url, params=None, timeout=None: Missing()

        with pytest.raises(lw.CalendarFeedError, match="older than this script"):
            client.get_sections("202710")

    def test_an_html_error_page_is_not_parsed_as_a_row(self, client):
        client.session.get = lambda url, params=None, timeout=None: FakeResponse(
            "<html>502 Bad Gateway</html>", content_type="text/html"
        )

        with pytest.raises(lw.CalendarFeedError, match="Expected CSV"):
            client.get_sections("202710")


class TestOutputRows:
    def test_section_columns_keep_their_order(self):
        row = lw.to_output_row({}, lw.SECTION_COLUMN_LABELS)

        assert list(row) == list(lw.SECTION_COLUMN_LABELS.values())
        assert list(row)[:3] == ["Term", "CRN", "Subject"]

    def test_numbers_stay_numbers(self):
        feed_row = next(iter(csv.DictReader(SECTION_CSV.splitlines())))

        row = lw.to_output_row(feed_row, lw.SECTION_COLUMN_LABELS)

        assert row["CRN"] == 16036
        assert row["Credit Hours"] == 4
        assert row["Enrollment Max"] == 30
        assert row["Enrollment Current"] == 25
        assert row["Room Capacity"] == 40
        assert row["Meeting Count"] == 2

    def test_a_missing_number_stays_blank_rather_than_zero(self):
        rows = list(csv.DictReader(SECTION_CSV.splitlines()))

        row = lw.to_output_row(rows[1], lw.SECTION_COLUMN_LABELS)

        assert row["Room Capacity"] == ""

    def test_a_number_column_that_is_not_a_number_is_left_alone(self):
        row = lw.to_output_row({"credit_hours": "1.5"}, lw.SECTION_COLUMN_LABELS)

        assert row["Credit Hours"] == "1.5"

    def test_day_codes_are_passed_through_untouched(self):
        feed_row = next(iter(csv.DictReader(SECTION_CSV.splitlines())))

        row = lw.to_output_row(feed_row, lw.SECTION_COLUMN_LABELS)

        assert row["Meeting Days"] == "TR"
        assert row["Meeting Times"] == "10:00-11:45"
        assert row["Location"] == "WENTW 214"


class TestFetchCourses:
    def _run(self, monkeypatch, tmp_path, **kwargs):
        feed = lw.CalendarFeedClient(base_url="https://calendar.example")
        feed.session = FakeSession({
            "/reports/sections": SECTION_CSV,
            "/reports/meeting_times": MEETING_CSV,
        })
        monkeypatch.setattr(lw, "CalendarFeedClient", lambda *a, **kw: feed)

        out = tmp_path / "out.csv"
        lw.fetch_courses("202710", str(out), "csv", verbose=False, **kwargs)
        return feed, out

    def test_one_row_per_section_by_default(self, monkeypatch, tmp_path):
        feed, out = self._run(monkeypatch, tmp_path)

        rows = read_csv(out)
        assert feed.session.calls[0][0].endswith("/reports/sections")
        assert len(rows) == 2
        assert [r["CRN"] for r in rows] == ["16036", "17309"]
        assert rows[0]["Meeting Days"] == "TR"
        assert "Day" not in rows[0]

    def test_by_meeting_gives_one_row_per_meeting_day(self, monkeypatch, tmp_path):
        feed, out = self._run(monkeypatch, tmp_path, by_meeting=True)

        rows = read_csv(out)
        assert feed.session.calls[0][0].endswith("/reports/meeting_times")
        assert len(rows) == 2
        assert [r["Day"] for r in rows] == ["tuesday", "thursday"]
        assert {r["CRN"] for r in rows} == {"16036"}
        assert "Meeting Days" not in rows[0]

    def test_json_names_the_grain_it_wrote(self, monkeypatch, tmp_path):
        feed = lw.CalendarFeedClient(base_url="https://calendar.example")
        feed.session = FakeSession({
            "/reports/sections": SECTION_CSV,
            "/reports/meeting_times": MEETING_CSV,
        })
        monkeypatch.setattr(lw, "CalendarFeedClient", lambda *a, **kw: feed)

        out = tmp_path / "out.json"
        lw.fetch_courses("202710", str(out), "json", verbose=False)
        assert "sections" in json.loads(out.read_text())

        lw.fetch_courses("202710", str(out), "json", verbose=False, by_meeting=True)
        assert "meeting_times" in json.loads(out.read_text())

    def test_an_unreachable_feed_exits_nonzero(self, monkeypatch, tmp_path):
        feed = lw.CalendarFeedClient(base_url="https://calendar.example")
        feed.session = FakeSession({})
        monkeypatch.setattr(lw, "CalendarFeedClient", lambda *a, **kw: feed)

        def explode(*args, **kwargs):
            raise lw.requests.RequestException("connection refused")

        feed.session.get = explode

        with pytest.raises(SystemExit) as exit_info:
            lw.fetch_courses("202710", str(tmp_path / "out.csv"), "csv", verbose=False)

        assert exit_info.value.code == 1


class TestExcel:
    def test_headers_match_the_chosen_grain(self, tmp_path):
        rows = [lw.to_output_row(r, lw.SECTION_COLUMN_LABELS)
                for r in csv.DictReader(SECTION_CSV.splitlines())]
        out = tmp_path / "out.xlsx"

        lw.save_as_excel(rows, "202710", str(out), lw.SECTION_COLUMN_LABELS, verbose=False)

        from openpyxl import load_workbook
        ws = load_workbook(out).active
        assert [c.value for c in ws[1]] == list(lw.SECTION_COLUMN_LABELS.values())
        assert ws.max_row == 3

    def test_numeric_cells_stay_numeric(self, tmp_path):
        rows = [lw.to_output_row(r, lw.SECTION_COLUMN_LABELS)
                for r in csv.DictReader(SECTION_CSV.splitlines())]
        out = tmp_path / "out.xlsx"

        lw.save_as_excel(rows, "202710", str(out), lw.SECTION_COLUMN_LABELS, verbose=False)

        from openpyxl import load_workbook
        ws = load_workbook(out).active
        crn_column = list(lw.SECTION_COLUMN_LABELS.values()).index("CRN") + 1
        assert ws.cell(row=2, column=crn_column).value == 16036
